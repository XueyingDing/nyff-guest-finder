# Step 3 - add Chinese actor names and film titles.
# Setup: pip install requests openpyxl
# Usage: run this in the folder that contains nyff64_films.xlsx (the output of step 2).
# It reads the "Very Famous Actors & Work" and "Famous Actors & Work" columns, looks up the
# Chinese names of the actors and the Chinese titles of their works on Wikidata, and writes a
# NEW workbook with a Chinese column inserted right after each of those two columns.
# Your original file is not modified.
import copy
import json
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _wikidata_common import (
    LookupUnavailable, claim_ids, count_wikis, get_json as _get_json, make_stdout_unicode_safe,
)

# This script prints Chinese names/titles to the console; see make_stdout_unicode_safe()'s
# docstring for why an uncaught encoding error here is worse than just a lost log line.
make_stdout_unicode_safe()

IN_FILE = "nyff64_films.xlsx"
OUT_FILE = "nyff64_films_with_chinese.xlsx"
CACHE_FILE = "zh_lookup_cache.json"          # remembers lookups so reruns are fast; safe to delete
OLD_CACHE_FILE = "wikidata_cache_v2.json"    # from the guest-lookup script; reused to get the same person

SOURCE_HEADERS = ["Very Famous Actors & Work", "Famous Actors & Work"]
NEW_HEADER_SUFFIX = " (Chinese)"
VERBOSE_RETRIES = True    # Print a line for every retry/wait (attempt count, reason, wait time)
WORKERS = 3   # parallel lookups; too many risks Wikidata rate limiting (429)

# Preference order for Chinese labels: Simplified first, then other variants
ZH_LANGS = ["zh-hans", "zh-cn", "zh", "zh-sg", "zh-hk", "zh-tw", "zh-hant"]

API = "https://www.wikidata.org/w/api.php"
SPARQL = "https://query.wikidata.org/sparql"
HTTP_HEADERS = {"User-Agent": "nyff-chinese-names/1.0 (personal use; python-requests)"}

# Film-related occupations: film director, actor, film actor, screenwriter, film producer, cinematographer
FILM_OCC = {"Q2526255", "Q33999", "Q10800557", "Q28389", "Q3282637", "Q222344"}


def get_json(url, params, retries=3):
    # Raises LookupUnavailable (a plain Exception subclass) if Wikidata never answers
    # successfully within `retries` tries -- it no longer silently returns {}. That used
    # to make a rate-limited (429) request look exactly like "person has no Chinese name /
    # is not in this actor's filmography", which then got cached as if it were true.
    return _get_json(url, params, HTTP_HEADERS, retries=retries, verbose=VERBOSE_RETRIES)


# ---------------------------------------------------------------- Wikidata helpers

def pick_zh(candidates):
    """candidates: list of (lang, text). Return the text in the most preferred Chinese variant."""
    by_lang = {}
    for lang, text in candidates:
        by_lang.setdefault(lang, text)
    for lang in ZH_LANGS:
        if lang in by_lang:
            return by_lang[lang]
    return ""


def search_qid(name):
    """Find the Wikidata id of a film person by English name ('' if not found).
    May raise LookupUnavailable if the request itself fails."""
    res = get_json(API, dict(action="wbsearchentities", search=name, language="en",
                             uselang="en", type="item", limit=10, format="json"))
    ids = [x["id"] for x in res.get("search", [])]
    if not ids:
        return ""
    ents = get_json(API, dict(action="wbgetentities", ids="|".join(ids),
                              props="labels|claims|sitelinks", languages="en",
                              format="json")).get("entities", {})
    best = None
    for qid in ids:
        e = ents.get(qid) or {}
        if "Q5" not in claim_ids(e, "P31"):                # humans only
            continue
        if not (claim_ids(e, "P106") & FILM_OCC):          # must be a film person
            continue
        label = e.get("labels", {}).get("en", {}).get("value", "")
        key = (label.lower() == name.lower(), count_wikis(e))
        if best is None or key > best[0]:
            best = (key, qid)
    return best[1] if best else ""


def fetch_zh_label(qid):
    """Look up a known Wikidata id's Chinese label ('' if it genuinely has none).
    Raises LookupUnavailable if the request fails -- callers must not cache that as if
    it were a real (possibly empty) label."""
    ents = get_json(API, dict(action="wbgetentities", ids=qid, props="labels",
                              languages="|".join(ZH_LANGS), format="json")).get("entities", {})
    labels = ents.get(qid, {}).get("labels", {})
    return pick_zh([(k, v["value"]) for k, v in labels.items()])


def find_actor(name, old_cache):
    """Return {'qid', 'zh', 'zh_checked'} for an actor's English name.

    zh_checked=True marks that the Chinese-label lookup genuinely completed (whether or
    not a label was found). A cache entry from before this field existed can't be told
    apart from a fresh "found nothing" result just by looking at qid/zh -- callers must
    treat a missing zh_checked as "needs re-verifying", not as "already known empty".

    Raises LookupUnavailable if a Wikidata request fails -- callers must not cache that
    as if it were a real (possibly empty) result. Because get_json() now raises instead
    of returning {} on repeated 429s, qid == "" here is trustworthy: it only happens when
    the search genuinely completed and found no matching film person."""
    qid = (old_cache.get(name) or {}).get("qid") or search_qid(name)
    if not qid:
        return {"qid": "", "zh": "", "zh_checked": True}
    return {"qid": qid, "zh": fetch_zh_label(qid), "zh_checked": True}


def sparql_literal(title):
    return '"' + title.replace("\\", "\\\\").replace('"', '\\"') + '"@en'


def query_films(qid, titles):
    """Among films this person directed or appeared in, match the given English titles.
    Returns {english_title: chinese_title_or_empty} for the titles that were found.
    Raises LookupUnavailable if the SPARQL request fails."""
    if not titles:
        return {}
    values = " ".join(sparql_literal(t) for t in titles)
    langs = ", ".join(f'"{lg}"' for lg in ZH_LANGS)
    q = f"""
    SELECT ?en ?zh ?zhLang WHERE {{
      VALUES ?en {{ {values} }}
      ?film rdfs:label ?en .
      {{ ?film wdt:P57 wd:{qid} . }} UNION {{ ?film wdt:P161 wd:{qid} . }}
      OPTIONAL {{
        ?film rdfs:label ?zh .
        FILTER(LANG(?zh) IN ({langs}))
      }}
      BIND(LANG(?zh) AS ?zhLang)
    }}
    """
    data = get_json(SPARQL, {"query": q, "format": "json"})
    collected = {}
    for b in data.get("results", {}).get("bindings", []):
        en = b["en"]["value"]
        collected.setdefault(en, [])
        if "zh" in b:
            collected[en].append((b["zhLang"]["value"], b["zh"]["value"]))
    return {en: pick_zh(cands) for en, cands in collected.items()}


# ---------------------------------------------------------------- Cell parsing / translating

def parse_cell(value):
    """'Adam Driver: Marriage Story, Paterson\\nScarlett Johansson: Her' -> [(actor, [titles]), ...]"""
    text = "" if value is None else str(value)
    # One actor per line; also tolerate the older '; ' separated format
    parts = re.split(r"\n|;\s+(?=[A-Z][^:;\n]*:\s)", text)
    out = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if ": " in part:
            actor, rest = part.split(": ", 1)
            titles = [t.strip() for t in rest.split(", ") if t.strip()]
            if titles == ["(none found)"] or (titles and titles[0].startswith("(works lookup incomplete")):
                titles = []
        else:
            actor, titles = part, []
        out.append((actor.strip(), titles))
    return out


def candidates_for(titles):
    """Every single title plus every adjacent pair joined by ', ' (some titles contain a comma,
    e.g. '... Deathly Hallows, Part 1'). Only titles that really exist for the actor will match."""
    cands = list(titles)
    cands += [titles[i] + ", " + titles[i + 1] for i in range(len(titles) - 1)]
    return cands


def translate_line(actor, titles, cache):
    zh_actor = cache["actors"].get(actor, {}).get("zh") or actor
    if not titles:
        return zh_actor

    def hit(t):
        return cache["works"].get(f"{actor}||{t}", {})

    out, i = [], 0
    while i < len(titles):
        if i + 1 < len(titles):
            merged = titles[i] + ", " + titles[i + 1]
            if hit(merged).get("found"):            # a single title that contains a comma
                out.append(hit(merged).get("zh") or merged)
                i += 2
                continue
        out.append(hit(titles[i]).get("zh") or titles[i])   # fall back to the English title
        i += 1
    return f"{zh_actor}: {', '.join(out)}"


# ---------------------------------------------------------------- Main

def main():
    wb = load_workbook(IN_FILE, rich_text=True)   # rich_text keeps the colored Guests column
    ws = wb.active

    headers = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value is not None}
    src_cols = []
    for h in SOURCE_HEADERS:
        if h.lower() not in headers:
            raise SystemExit(f'Column "{h}" not found. Headers in the file: '
                             f"{[c.value for c in ws[1]]}")
        src_cols.append(headers[h.lower()])
    src_cols.sort()

    # ---- collect everything that needs a lookup
    need = {}   # actor -> set of candidate titles
    for col in src_cols:
        for r in range(2, ws.max_row + 1):
            for actor, titles in parse_cell(ws.cell(row=r, column=col).value):
                need.setdefault(actor, set()).update(candidates_for(titles))
    print(f"{len(need)} distinct actors to translate")

    cache = {"actors": {}, "works": {}}
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, encoding="utf-8") as f:
            cache = json.load(f)
    old_cache = {}
    if os.path.exists(OLD_CACHE_FILE):
        with open(OLD_CACHE_FILE, encoding="utf-8") as f:
            old_cache = json.load(f)

    def work(actor):
        """Look up an actor's Chinese name (if not already cached) and any of their
        titles' Chinese names that are still missing. The two lookups are allowed to
        fail independently: if the actor's identity is already known (or just got
        resolved) but the works/title SPARQL query fails, we still return that
        identity so it gets cached -- only the missing titles stay missing and get
        retried next run, instead of the actor's already-successful info being
        thrown away along with the unrelated failure."""
        info = cache["actors"].get(actor)
        if info is not None and info.get("qid") and not info.get("zh_checked"):
            # A cache entry from before zh_checked existed: the qid is trustworthy (it
            # came from a real successful search), but an empty zh here is ambiguous --
            # it might be genuine, or it might be a 429 during the label lookup that an
            # older version of this script silently cached as empty. Re-fetch just the
            # label for the already-known qid; no need to redo the identity search.
            info = {"qid": info["qid"], "zh": fetch_zh_label(info["qid"]), "zh_checked": True}
        elif info is None:
            info = find_actor(actor, old_cache)   # may raise LookupUnavailable -> propagates to safe_work
        missing = [t for t in need[actor] if f"{actor}||{t}" not in cache["works"]]
        updates = {}
        works_err = ""
        if missing and info.get("qid"):
            try:
                found = query_films(info["qid"], missing)
                for t in missing:
                    updates[f"{actor}||{t}"] = {"found": t in found, "zh": found.get(t, "")}
            except Exception as ex:
                works_err = str(ex)[:100]
        return actor, info, updates, works_err

    def safe_work(actor):
        print(f"  [start] {actor}", flush=True)
        try:
            return work(actor)
        except LookupUnavailable as ex:
            print(f"[{actor}] lookup unavailable: {str(ex)[:100]}", flush=True)
            return actor, None, {}, str(ex)[:100]
        except Exception as ex:
            print(f"[{actor}] lookup failed: {str(ex)[:100]}", flush=True)
            return actor, None, {}, str(ex)[:100]

    todo = list(need)
    t0 = time.time()
    resolved, failed_actors, works_incomplete = [], [], []
    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futures = [pool.submit(safe_work, a) for a in todo]
        for done, fut in enumerate(as_completed(futures), 1):
            try:
                actor, info, updates, werr = fut.result()
                if info is not None:
                    cache["actors"][actor] = info
                    cache["works"].update(updates)
                    with open(CACHE_FILE, "w", encoding="utf-8") as f:
                        json.dump(cache, f, ensure_ascii=False, indent=1)
                    resolved.append(actor)
                    if werr:
                        works_incomplete.append(actor)
                else:
                    failed_actors.append(actor)
                zh = (info or {}).get("zh", "")
                status = "ok" if info is not None else "FAILED (will retry next run)"
                extra = f"; works/titles incomplete: {werr} (will retry next run)" if werr else ""
                print(f"[{done}/{len(todo)}] {actor} -> {zh or '(no Chinese name found)'} [{status}{extra}]",
                      flush=True)
            except Exception as ex:
                # Never let one actor's unexpected failure abort the whole batch and
                # strand every other in-flight lookup with nothing saved.
                print(f"[{done}/{len(todo)}] unexpected error recording a result: {ex!r}", flush=True)
    print(f"Lookups took {time.time() - t0:.0f} s", flush=True)
    print(f"\nLookup summary: {len(resolved)} actors resolved, {len(failed_actors)} failed this run "
          f"(will retry next run), {len(works_incomplete)} had an incomplete works/title lookup "
          f"(will retry next run)")
    if failed_actors:
        print("  Failed this run: " + ", ".join(failed_actors))
    if works_incomplete:
        print("  Works/titles incomplete: " + ", ".join(works_incomplete))

    # ---- insert the new columns right after each source column (right to left keeps indexes valid)
    old_widths = {c: ws.column_dimensions[get_column_letter(c)].width
                  for c in range(1, ws.max_column + 1)}
    for col in reversed(src_cols):
        ws.insert_cols(col + 1)

    def new_pos(old_col):   # final position of an original column
        return old_col + sum(1 for s in src_cols if s < old_col)

    for old_col, width in old_widths.items():
        if width:
            ws.column_dimensions[get_column_letter(new_pos(old_col))].width = width
    for s in src_cols:
        src_final = new_pos(s)
        ws.column_dimensions[get_column_letter(src_final + 1)].width = old_widths.get(s) or 40

    # ---- fill in the Chinese columns, copying the formatting (color, bold, wrapping) of the source
    for s in src_cols:
        src_final = new_pos(s)
        dst = src_final + 1
        head_src = ws.cell(row=1, column=src_final)
        head_dst = ws.cell(row=1, column=dst)
        head_dst.value = f"{head_src.value}{NEW_HEADER_SUFFIX}"
        head_dst._style = copy.copy(head_src._style)
        for r in range(2, ws.max_row + 1):
            src_cell = ws.cell(row=r, column=src_final)
            lines = [translate_line(actor, titles, cache)
                     for actor, titles in parse_cell(src_cell.value)]
            dst_cell = ws.cell(row=r, column=dst)
            dst_cell.value = "\n".join(lines)
            dst_cell._style = copy.copy(src_cell._style)

    wb.save(OUT_FILE)
    print(f"\nSaved: {os.path.abspath(OUT_FILE)}")


if __name__ == "__main__":
    main()
